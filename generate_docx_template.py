"""
Generate jharkhand_master_register.docx entirely from hard-coded data.

Usage:
    pip install python-docx
    python generate_docx_template.py
Output:
    jharkhand_master_register.docx  (created in the current directory)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ---------------------------------------------------------------------------
# Hard-coded data
# ---------------------------------------------------------------------------

GENERATED_ON = "15-05-2026 18:29"
TOTAL_RECORDS = 14

COVER_FILES = [
    {
        "sr": 1,
        "file_no": "14/वि.रु.भि.भ 784/25",
        "sanchika_no": "15/ग.नि.श. – 83/2025",
        "name_hi": "श्री शशिकान्त कुबे",
        "name_en": "Sri Shashikant Kube",
        "rank_id": "आरक्षी – 620 (Aarakshi-620)",
        "unit": "Jha.Sa.Pu.-07, Hazaribagh",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Shashikant Kube, Aarakshi-620, Jha.Sa.Pu.-07, Hazaribagh."
        ),
        "department": "Home, Prison & Disaster Management Dept.",
        "year": "2025",
    },
    {
        "sr": 2,
        "file_no": "14/वि.रा.नि.अ. – 412/26",
        "sanchika_no": "15/ग.नि.उ. – 129/2024",
        "name_hi": "साबिहा रबातुन",
        "name_en": "Sabiha Rabatun",
        "rank_id": "म. आरक्षी – 131 (M. Aarakshi-131)",
        "unit": "Bokaro District Force",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sabiha Rabatun, M. Aarakshi-131, Bokaro District Force."
        ),
        "department": "Home, Prison & Disaster Management Dept.",
        "year": "2024",
    },
    {
        "sr": 3,
        "file_no": "13/वि.रा.नि.अ. – 360/26",
        "sanchika_no": "11/गृ.स्था. – गृह निर्माण अग्रिम – 12/2025",
        "name_hi": "श्रीमती संध्या खलखो",
        "name_en": "Smt. Sandhya Khalkho",
        "rank_id": "प्रशाखा पदाधिकारी (Branch Officer)",
        "unit": "Forest Env. & Climate Change Dept., Ranchi",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Smt. Sandhya Khalkho, Branch Officer, Forest Environment & "
            "Climate Change Department, Jharkhand, Ranchi."
        ),
        "department": "Forest, Environment & Climate Change Dept.",
        "year": "2025",
    },
    {
        "sr": 4,
        "file_no": "14/वि.उठ.नि.अ. – 35/26",
        "sanchika_no": "16/गृह अग्रिम – 327/2025",
        "name_hi": "श्री निरंजन प्रसाद साह",
        "name_en": "Sri Niranjan Prasad Sah",
        "rank_id": "आरक्षी – 1562 (Aarakshi-1562)",
        "unit": "Jharkhand Jaguar (STF), Ranchi",
        "subject_en": (
            "Regarding approval of advance amount for home construction for "
            "Sri Niranjan Prasad Sah, Aarakshi-1562, Jharkhand Jaguar (STF), Ranchi."
        ),
        "department": "Home, Prison & Disaster Management Dept.",
        "year": "2025",
    },
    {
        "sr": 5,
        "file_no": "14/वि.च्र.नि.अ. – 33/26",
        "sanchika_no": "16/गृह अग्रिम – 321/2025",
        "name_hi": "श्री सीताराम रविदास",
        "name_en": "Sri Sitaram Ravidas",
        "rank_id": "आ. – 605 (A.-605)",
        "unit": "Jharkhand Jaguar",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Sitaram Ravidas, A.-605, Jharkhand Jaguar."
        ),
        "department": "Home, Prison & Disaster Management Dept.",
        "year": "2025",
    },
]

REGISTER_ENTRIES = [
    {
        "page": 299,
        "sr": 393,
        "file_no": "14/वि.च्र.नि.अ. – 393/26",
        "name_hi": "ज्ञानिमान साह",
        "name_en": "Gyaniman Sah",
        "rank_id": "आ. – 400 (A.-400)",
        "unit_hi": "जामताड़ा जिलाबल",
        "unit_en": "Jamtara District Force",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Gyaniman Sah, A.-400, Jamtara District Force."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 394,
        "file_no": "14/वि.क्रि.ग्रु.नि.भ. – 394/26",
        "name_hi": "श्री राहुल अजीरिमा",
        "name_en": "Sri Rahul Ajirimaa",
        "rank_id": "आरक्षी – 438 / अति. आरक्षी – 02 (Aarakshi-438 / Addl. Aarakshi-02)",
        "unit_hi": "बहुई",
        "unit_en": "Bahui",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Rahul Ajirimaa, Aarakshi-438, Addl. Aarakshi-02, Bahui."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 395,
        "file_no": "14/वि.गु.नि.अ. – 395/26",
        "name_hi": "श्री मनोज कुमार प्रजापति",
        "name_en": "Sri Manoj Kumar Prajapati",
        "rank_id": "पा.आ. / 2 (PA/2)",
        "unit_hi": "JAP-07, H2",
        "unit_en": "JAP-07, H2",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Manoj Kumar Prajapati, PA/2, JAP-07, H2."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 396,
        "file_no": "14/वि.च्र.नि.अ. – 396/26",
        "name_hi": "श्री अमित कुमार",
        "name_en": "Sri Amit Kumar",
        "rank_id": "सिपाही – 653 (Sipahi-653)",
        "unit_hi": "JAP-03, गोविंदपुर, धनबाद",
        "unit_en": "JAP-03, Govindpur, Dhanbad",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Amit Kumar, Sipahi-653, JAP-03, Govindpur, Dhanbad."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 397,
        "file_no": "14/वि.च्र.नि.अ. – 397/26",
        "name_hi": "श्री महेन्द्र यादव",
        "name_en": "Sri Mahendra Yadav",
        "rank_id": "चालक / हवलदार (Driver/Havaldar)",
        "unit_hi": "JAP-7, H2B",
        "unit_en": "JAP-7, H2B",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Mahendra Yadav, Driver/Havaldar, JAP-7, H2B."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 398,
        "file_no": "14/वि.च्र.नि.अ. – 398/26",
        "name_hi": "श्री लेखोट कुमार",
        "name_en": "Sri Lekhot Kumar",
        "rank_id": "आ. – 94 (A.-94)",
        "unit_hi": "JAP-6, धनबाद",
        "unit_en": "JAP-6, Dhanbad",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Lekhot Kumar, A.-94, JAP-6, Dhanbad."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 399,
        "file_no": "14/वि.च्र.नि.अ. – 399/26",
        "name_hi": "श्री सुनील कुमार वर्मा",
        "name_en": "Sri Sunil Kumar Verma",
        "rank_id": "चाव. छाप. – 3559 (No.-3559)",
        "unit_hi": "रांची जिलाबल",
        "unit_en": "Ranchi District Force",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Sunil Kumar Verma, No.-3559, Ranchi District Force."
        ),
        "remarks": "–",
    },
    {
        "page": 299,
        "sr": 400,
        "file_no": "14/वि.च्र.नि.अ. – 400/26",
        "name_hi": "श्री सुमित कुमार अंजेजा सह एक अन्य",
        "name_en": "Sri Sumit Kumar Anjeja & One Other",
        "rank_id": "–",
        "unit_hi": "संचार एवं तकनीकी सेवाएं, रांची",
        "unit_en": "Communication & Technical Services, Ranchi",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Sumit Kumar Anjeja and one other, "
            "Communication & Technical Services, Ranchi."
        ),
        "remarks": "सह एक अन्य (Along with one other)",
    },
    {
        "page": 299,
        "sr": 401,
        "file_no": "14/वि.च्र.नि.अ. – 401/26",
        "name_hi": "श्री धरिंद्र विश्वकर्मा",
        "name_en": "Sri Dharindra Vishwakarma",
        "rank_id": "आ. – 33 (A.-33)",
        "unit_hi": "JAP-03, गोविंदपुर, धनबाद",
        "unit_en": "JAP-03, Govindpur, Dhanbad",
        "subject_en": (
            "Regarding approval of home construction advance for "
            "Sri Dharindra Vishwakarma, A.-33, JAP-03, Govindpur, Dhanbad."
        ),
        "remarks": "–",
    },
]

SUMMARY = [
    ("Total Records", "14"),
    ("Cover Files", "5"),
    ("Register Entries", "9"),
    ("Register Page", "299  (Sr. 393 – 401)"),
    ("State", "Jharkhand (झारखण्ड)"),
    ("Departments", "1. Home, Prison & Disaster Mgmt.\n2. Forest, Environment & Climate Change"),
    (
        "Common Purpose",
        "Home Construction Advance Approval\n(गृह निर्माण अग्रिम स्वीकृति)",
    ),
    (
        "Units / Locations",
        "Hazaribagh · Bokaro · Ranchi · Jamtara · Dhanbad · Bahui\n"
        "JAP-03 · JAP-06 · JAP-07 · STF · H2 · H2B",
    ),
    ("File Series", "14/वि.च्र.नि.अ. – XXX/26"),
    ("Database", "govt_documents.db  (SQLite, 19 columns)"),
    ("Generated On", "15-05-2026 18:29:43"),
    ("Legend – row colour", "Green tint = Cover File  |  Yellow tint = Register Entry"),
]

# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------

C_HEADER_BG   = RGBColor(0x1F, 0x49, 0x7D)   # dark navy
C_HEADER_FG   = RGBColor(0xFF, 0xFF, 0xFF)   # white
C_COVER_BG    = RGBColor(0xE2, 0xEF, 0xDA)   # light green tint
C_REGISTER_BG = RGBColor(0xFF, 0xFF, 0xCC)   # light yellow tint
C_SECTION_BG  = RGBColor(0xDD, 0xEB, 0xF7)   # light blue for section headings
C_SUMMARY_HDR = RGBColor(0x2F, 0x75, 0xB6)   # mid blue

# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def _rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_cell_bg(cell, color: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  _rgb_hex(color))
    tcPr.append(shd)


def set_cell_borders(cell, border_size: int = 4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"),   "single")
        b.set(qn("w:sz"),    str(border_size))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "4F81BD")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def para_in_cell(cell, text: str, bold=False, italic=False,
                 size_pt=9, color: RGBColor = None,
                 align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0].clear()
    p   = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return p


def add_section_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"  {text}  ")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = C_HEADER_FG
    # shade the paragraph via direct XML (paragraph shading)
    pPr  = p._p.get_or_add_pPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  _rgb_hex(C_HEADER_BG))
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(4)
    return p


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # --- Page margins ---
    section = doc.sections[0]
    section.page_width   = Cm(29.7)
    section.page_height  = Cm(21.0)
    section.orientation  = 1                  # landscape
    section.left_margin  = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin   = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    # --- Default paragraph style ---
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # =========================================================
    # TITLE BLOCK
    # =========================================================
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after  = Pt(2)
    r = title.add_run("झारखण्ड सरकार — Master Document Register")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = C_HEADER_BG

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after  = Pt(2)
    r2 = sub.add_run("Government of Jharkhand  |  गृह निर्माण अग्रिम स्वीकृति")
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.color.rgb = C_SUMMARY_HDR

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after  = Pt(6)
    r3 = meta.add_run(
        f"Generated: {GENERATED_ON}   |   Total Records: {TOTAL_RECORDS}   |   "
        "Dept: Home, Prison & Disaster Mgmt. + Forest & Climate Change"
    )
    r3.italic = True
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # =========================================================
    # SECTION 1 — ALL RECORDS (master table)
    # =========================================================
    add_section_heading(doc, "Section 1 — All Records  (Master Register)")

    MASTER_COLS = [
        "#", "Source\nType", "Reg.\nPage", "Reg. Sr.\nNo.",
        "File No.", "Sanchika No.", "Name (Hindi)", "Name (English)",
        "Rank / ID", "Unit (English)", "Subject (English)", "Purpose", "Remarks",
    ]
    MASTER_WIDTHS = [0.8, 1.5, 1.2, 1.2, 3.2, 3.5, 2.8, 2.8, 3.0, 3.2, 5.5, 3.5, 2.2]

    tbl = doc.add_table(rows=1, cols=len(MASTER_COLS))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, col in enumerate(MASTER_COLS):
        set_cell_bg(hdr_cells[i], C_HEADER_BG)
        set_cell_borders(hdr_cells[i])
        para_in_cell(hdr_cells[i], col, bold=True, size_pt=8,
                     color=C_HEADER_FG, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Cover file rows (green tint)
    for cf in COVER_FILES:
        row = tbl.add_row().cells
        bg  = C_COVER_BG
        vals = [
            str(cf["sr"]), "Cover File", "–", "–",
            cf["file_no"], cf["sanchika_no"],
            cf["name_hi"], cf["name_en"], cf["rank_id"], cf["unit"],
            cf["subject_en"], "Home Construction Advance Approval", "–",
        ]
        for i, v in enumerate(vals):
            set_cell_bg(row[i], bg)
            set_cell_borders(row[i])
            para_in_cell(row[i], v, size_pt=8)

    # Register entry rows (yellow tint)
    for re_ in REGISTER_ENTRIES:
        row = tbl.add_row().cells
        bg  = C_REGISTER_BG
        vals = [
            str(COVER_FILES[-1]["sr"] + (re_["sr"] - REGISTER_ENTRIES[0]["sr"] + 1)),
            "Register",
            str(re_["page"]),
            str(re_["sr"]),
            re_["file_no"], "–",
            re_["name_hi"], re_["name_en"], re_["rank_id"], re_["unit_en"],
            re_["subject_en"], "Home Construction Advance Approval", re_["remarks"],
        ]
        for i, v in enumerate(vals):
            set_cell_bg(row[i], bg)
            set_cell_borders(row[i])
            para_in_cell(row[i], v, size_pt=8)

    set_col_widths(tbl, MASTER_WIDTHS)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================
    # SECTION 2 — COVER FILES
    # =========================================================
    add_section_heading(doc, "Section 2 — Cover Files")

    CF_COLS = [
        "#", "File No.", "Sanchika No.", "Name (Hindi)", "Name (English)",
        "Rank / ID", "Unit", "Subject (English)", "Department", "Year",
    ]
    CF_WIDTHS = [0.8, 3.2, 3.5, 2.8, 2.8, 3.0, 3.2, 5.5, 3.8, 1.2]

    tbl2 = doc.add_table(rows=1, cols=len(CF_COLS))
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl2.style = "Table Grid"

    hdr2 = tbl2.rows[0].cells
    for i, col in enumerate(CF_COLS):
        set_cell_bg(hdr2[i], C_HEADER_BG)
        set_cell_borders(hdr2[i])
        para_in_cell(hdr2[i], col, bold=True, size_pt=8,
                     color=C_HEADER_FG, align=WD_ALIGN_PARAGRAPH.CENTER)

    for cf in COVER_FILES:
        row = tbl2.add_row().cells
        vals = [
            str(cf["sr"]), cf["file_no"], cf["sanchika_no"],
            cf["name_hi"], cf["name_en"], cf["rank_id"], cf["unit"],
            cf["subject_en"], cf["department"], cf["year"],
        ]
        for i, v in enumerate(vals):
            set_cell_bg(row[i], C_COVER_BG)
            set_cell_borders(row[i])
            para_in_cell(row[i], v, size_pt=8)

    set_col_widths(tbl2, CF_WIDTHS)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================
    # SECTION 3 — REGISTER ENTRIES (Page 299)
    # =========================================================
    add_section_heading(doc, "Section 3 — Register Entries  (Page 299, Sr. 393–401)")

    RE_COLS = [
        "Page", "Sr. No.", "File No.", "Name (Hindi)", "Name (English)",
        "Rank / ID", "Unit (Hindi)", "Unit (English)",
        "Subject (English)", "Purpose", "Remarks",
    ]
    RE_WIDTHS = [1.0, 1.2, 3.2, 2.8, 2.8, 3.0, 2.8, 3.2, 5.5, 3.5, 2.2]

    tbl3 = doc.add_table(rows=1, cols=len(RE_COLS))
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl3.style = "Table Grid"

    hdr3 = tbl3.rows[0].cells
    for i, col in enumerate(RE_COLS):
        set_cell_bg(hdr3[i], C_HEADER_BG)
        set_cell_borders(hdr3[i])
        para_in_cell(hdr3[i], col, bold=True, size_pt=8,
                     color=C_HEADER_FG, align=WD_ALIGN_PARAGRAPH.CENTER)

    for re_ in REGISTER_ENTRIES:
        row = tbl3.add_row().cells
        vals = [
            str(re_["page"]), str(re_["sr"]), re_["file_no"],
            re_["name_hi"], re_["name_en"], re_["rank_id"],
            re_["unit_hi"], re_["unit_en"],
            re_["subject_en"],
            "Home Construction Advance Approval\n(गृह निर्माण अग्रिम स्वीकृति)",
            re_["remarks"],
        ]
        for i, v in enumerate(vals):
            set_cell_bg(row[i], C_REGISTER_BG)
            set_cell_borders(row[i])
            para_in_cell(row[i], v, size_pt=8)

    set_col_widths(tbl3, RE_WIDTHS)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================
    # SECTION 4 — SUMMARY
    # =========================================================
    add_section_heading(doc, "Section 4 — Summary")

    tbl4 = doc.add_table(rows=0, cols=2)
    tbl4.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl4.style = "Table Grid"

    for field, value in SUMMARY:
        row = tbl4.add_row().cells
        set_cell_bg(row[0], C_SUMMARY_HDR)
        set_cell_borders(row[0])
        para_in_cell(row[0], field, bold=True, size_pt=9, color=C_HEADER_FG)

        set_cell_borders(row[1])
        para_in_cell(row[1], value, size_pt=9)

    for row in tbl4.rows:
        row.cells[0].width = Cm(5)
        row.cells[1].width = Cm(12)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================
    # FOOTER NOTE
    # =========================================================
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(10)
    fr = footer_p.add_run(
        "Birsa Munda Kendriya Kara, Hotwar, Ranchi  |  "
        "Ph: 0651-2276060  |  Fax: 2276012"
    )
    fr.font.size = Pt(8)
    fr.italic = True
    fr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    return doc


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    out_path = "jharkhand_master_register.docx"
    doc = build_document()
    doc.save(out_path)
    print(f"Document saved: {out_path}")
